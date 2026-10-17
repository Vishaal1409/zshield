"""Build docs/deck/ZShield_deck_draft.pptx from docs/deck/storyline.md.

Run from the repo root:  .venv/bin/python docs/deck/build_deck.py
"""
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
STORY = HERE / "storyline.md"
OUT = HERE / "ZShield_deck_draft.pptx"

ACCENT = RGBColor(0x0F, 0x62, 0xFE)
DARK = RGBColor(0x16, 0x16, 0x16)
GREY = RGBColor(0x52, 0x52, 0x52)
LIGHT = RGBColor(0xE8, 0xE8, 0xE8)
BOX_FILL = RGBColor(0xF2, 0xF4, 0xF8)
FONT = "Arial"  # IBM Plex Sans may be missing on the viewing machine


def parse(text):
    slides, sources, current, key, section = [], [], None, None, None
    for raw in text.splitlines():
        line = raw.rstrip()
        if line.startswith("## "):
            section = "slide" if line.startswith("## Slide ") else (
                "sources" if line.startswith("## Sources") else "other")
            key = None
            if section == "slide":
                m = re.match(r"## Slide (\d+): (.*)", line)
                current = {"num": int(m.group(1)), "name": m.group(2), "bullets": []}
                slides.append(current)
            continue
        if line.strip() == "---":
            key = None
            continue
        if section == "sources" and line.startswith("- "):
            sources.append(line[2:])
            continue
        if section != "slide":
            continue
        m = re.match(r"\*\*(\w[\w ]*):\*\*\s*(.*)", line)
        if m:
            key = m.group(1).lower()
            if key == "bullets":
                continue
            current[key] = m.group(2)
        elif line.startswith("- ") and key == "bullets":
            current["bullets"].append(line[2:])
        elif line.strip() and key and key != "bullets":
            current[key] += " " + line.strip()
    return slides, sources


def add_text(slide, x, y, w, h, text, size, bold=False, color=DARK, italic=False,
             anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    f = r.font
    f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
    f.color.rgb = color
    return tb


def est_lines(text, size, width_in):
    chars = max(1, int(width_in * 72 / (size * 0.5)))
    return -(-len(text) // chars)


def footer(slide, n):
    add_text(slide, 0.5, 7.0, 8, 0.35, "ZShield | IBM Z Datathon 2026", 12, color=GREY)
    add_text(slide, 11.8, 7.0, 1.0, 0.35, str(n), 12, color=GREY, align=PP_ALIGN.RIGHT)


def base(prs, headline, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.25), Inches(7.5))
    bar.fill.solid(); bar.fill.fore_color.rgb = ACCENT; bar.line.fill.background()
    add_text(s, 0.55, 0.3, 12.3, 1.3, headline, 28, bold=True, color=DARK,
             anchor=MSO_ANCHOR.MIDDLE)
    ln = s.shapes.add_connector(1, Inches(0.6), Inches(1.65), Inches(12.7), Inches(1.65))
    ln.line.color.rgb = ACCENT; ln.line.width = Pt(2)
    footer(s, n)
    return s


def bullets(slide, items, x, y, w, h, size=20):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        r = p.add_run()
        r.text = "•  " + it
        r.font.name, r.font.size = FONT, Pt(size)
        r.font.color.rgb = DARK
    return tb


def placeholder(slide, text, x, y, w, h):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = BOX_FILL
    sh.line.color.rgb = GREY; sh.line.width = Pt(1.5)
    sh.line.dash_style = MSO_LINE.DASH
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.2)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name, r.font.size, r.font.bold = FONT, Pt(18), True
    r.font.color.rgb = GREY


def build():
    slides, sources = parse(STORY.read_text(encoding="utf-8"))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    warnings = []
    for d in slides:
        s = base(prs, d["headline"], d["num"])
        if est_lines(d["headline"], 28, 12.2) > 2:
            warnings.append(f"slide {d['num']}: headline may be 3 lines")
        has_pl = d["num"] == 1
        body_h = 3.0 if has_pl else 3.9
        bsize = 20
        def need(sz):
            return sum(est_lines("\u2022  " + x, sz, 6.7) for x in d["bullets"]) * (sz / 60 + 0.01) + 0.15 * len(d["bullets"])
        if need(20) > body_h:
            bsize = 18
        used = need(bsize)
        bullets(s, d["bullets"], 0.6, 1.85, 6.9, body_h, size=bsize)
        if used > body_h:
            warnings.append(f"slide {d['num']}: bullets may overflow ({used:.1f} in > {body_h})")
        placeholder(s, d["box"], 7.8, 1.9, 5.0, 2.8 if has_pl else 3.6)
        if has_pl:
            pl = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(4.85),
                                    Inches(12.2), Inches(1.25))
            pl.fill.solid(); pl.fill.fore_color.rgb = ACCENT; pl.line.fill.background()
            tf = pl.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = Inches(0.25)
            r = tf.paragraphs[0].add_run(); r.text = d["problem line"]
            r.font.name, r.font.size, r.font.bold = FONT, Pt(20), True
            r.font.color.rgb = RGBColor(255, 255, 255)
        small = d.get("small", "")
        sy = 6.2 if has_pl else 5.85
        sh_ = 0.75 if has_pl else 1.1
        add_text(s, 0.6, sy, 12.2, sh_, small, 12, color=GREY, italic=True)
        s.notes_slide.notes_text_frame.text = (
            d["notes"] + "\n\nVisual needed: " + d["visual"] + " (Supplier: " + d["supplier"] + ")")

    # Appendix: sources
    n = len(slides) + 1
    s = base(prs, "Appendix: sources and notes", n)
    tb = s.shapes.add_textbox(Inches(0.6), Inches(1.85), Inches(12.2), Inches(5.0))
    tf = tb.text_frame; tf.word_wrap = True
    for i, src in enumerate(sources):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        r = p.add_run(); r.text = "•  " + src
        r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(13), DARK
    s.notes_slide.notes_text_frame.text = "Appendix. All figures are from docs/problem/stats.md, quoted with their source."
    prs.core_properties.author = "ZShield team"
    prs.core_properties.title = "ZShield deck draft"
    prs.save(OUT)
    print("saved", OUT)
    for w in warnings:
        print("WARNING:", w)


if __name__ == "__main__":
    build()
